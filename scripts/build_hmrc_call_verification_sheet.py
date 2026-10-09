from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "correspondence" / "HMRC_LedgerHouse_call_verification_sheet.docx"

NAVY = "17365D"
BLUE = "2E75B6"
PALE_BLUE = "EAF2F8"
PALE_GOLD = "FFF4CE"
LIGHT_GREY = "F3F5F7"
MID_GREY = "667085"
WHITE = "FFFFFF"
BLACK = "111827"
RED = "9C0006"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=100, bottom=70, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, color="CDD5DF", size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def set_cell_width(cell, width_twips):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def set_table_widths(table, widths_twips):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_twips)))
    tbl_w.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_twips:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths_twips):
            set_cell_width(cell, width)


def style_run(run, size=9, bold=False, color=BLACK):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_cell_text(cell, text, *, bold=False, color=BLACK, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    style_run(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_compact_paragraph(container, text, *, size=9, bold=False, color=BLACK, align=None, before=0, after=0):
    p = container.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    style_run(run, size=size, bold=bold, color=color)
    return p


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(297)
    section.page_height = Mm(210)
    section.top_margin = Mm(10)
    section.bottom_margin = Mm(9)
    section.left_margin = Mm(11)
    section.right_margin = Mm(11)
    section.header_distance = Mm(5)
    section.footer_distance = Mm(5)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(9)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(0)

    header = section.header
    hp = header.paragraphs[0]
    hp.text = "KAFS LIMITED | LEDGERHOUSE | HMRC DEVELOPER HUB"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_run(hp.runs[0], size=7.5, bold=True, color=MID_GREY)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(1)
    run = title.add_run("HMRC ACCOUNT VERIFICATION AND FOLLOW-UP CALL SHEET")
    style_run(run, size=17, bold=True, color=NAVY)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(5)
    style_run(
        subtitle.add_run("Purpose: verify the Developer Hub account and obtain an escalation outcome for the overdue LedgerHouse VAT MTD production-credentials request."),
        size=8.7,
        color=MID_GREY,
    )

    warning = doc.add_table(rows=1, cols=1)
    warning.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_widths(warning, [15550])
    set_table_borders(warning, color="D6B656", size="5")
    wc = warning.cell(0, 0)
    set_cell_shading(wc, PALE_GOLD)
    add_cell_text(
        wc,
        "SECURITY: Use HMRC's official number. Never disclose a password, MFA/OTP code, client secret, OAuth token, full card details or online-banking credentials. Give the full UTR only if the verified HMRC adviser specifically asks for it.",
        bold=True,
        color=RED,
        size=8.3,
    )

    add_compact_paragraph(doc, "Prepared answers", size=10.5, bold=True, color=BLUE, before=5, after=2)

    rows = [
        ("What is your full name?", "Godfred Frimpong"),
        ("What is your role?", "Software Development Officer and responsible individual for the LedgerHouse Developer Hub application."),
        ("What organisation do you represent?", "KAFS Limited, trading as Kafs Automation."),
        ("What is your business address?", "5 Brayford Square, London E1 0SG."),
        ("What email is registered/used?", "info@kafs-ltd.com"),
        ("What telephone numbers can you confirm?", "Mobile: +44 7939 823988 | Landline: +44 20 8434 9531"),
        ("What is the software/application name?", "LedgerHouse."),
        ("What is the organisation/software URL?", "https://ledgerhouse.finaccord.pro"),
        ("Which API access was requested?", "VAT (Making Tax Digital) production credentials."),
        ("What is the request reference?", "2026-VXY565."),
        ("When was it acknowledged?", "13 July 2026 at 10:39 pm (Europe/London)."),
        ("What does the portal currently show?", "Credentials requested. VAT (MTD) is the subscribed production API."),
        ("Can you confirm the Corporation Tax UTR?", "Yes. Read it from the secure HMRC record only if asked; it is intentionally not printed on this portable sheet."),
        ("Is there another recent support reference?", "Yes: 2026-NLG275 is for the separate 2-step verification matter, not the production-credentials request."),
        ("Can you confirm the application ID?", "Open Developer Hub > Applications > LedgerHouse and read it only if the adviser asks. Do not volunteer client credentials."),
    ]

    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_widths(table, [4550, 11000])
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for cell, text in zip(hdr.cells, ("Possible verification question", "Prepared answer")):
        set_cell_shading(cell, NAVY)
        add_cell_text(cell, text, bold=True, color=WHITE, size=9)

    for index, (question, answer) in enumerate(rows):
        cells = table.add_row().cells
        set_table_widths(table, [4550, 11000])
        if index % 2:
            set_cell_shading(cells[0], LIGHT_GREY)
            set_cell_shading(cells[1], LIGHT_GREY)
        add_cell_text(cells[0], question, bold=True, color=NAVY, size=8.25)
        add_cell_text(cells[1], answer, size=8.25)

    add_compact_paragraph(doc, "Questions to ask before ending the call", size=10.5, bold=True, color=BLUE, before=5, after=2)
    questions = [
        "Is reference 2026-VXY565 still open and assigned to the correct production-credentials team?",
        "Has the responsible-person verification email been sent, received and cleared?",
        "Is any evidence, demonstration, correction or action outstanding from KAFS Limited?",
        "Please escalate the overdue request and give me the escalation/case reference.",
        "What is the expected response date, and which email address will receive the decision?",
        "Please add a note that the Developer Hub still shows 'Credentials requested'.",
    ]
    qtable = doc.add_table(rows=3, cols=2)
    qtable.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_widths(qtable, [7775, 7775])
    set_table_borders(qtable, color="D9E2F3", size="4")
    for idx, question in enumerate(questions):
        cell = qtable.cell(idx % 3, idx // 3)
        set_cell_shading(cell, PALE_BLUE)
        add_cell_text(cell, f"{idx + 1}. {question}", size=8.15)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(
        fp.add_run("Call record: adviser __________________  date/time __________________  escalation reference __________________  promised response date __________________"),
        size=7.6,
        bold=True,
        color=MID_GREY,
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
