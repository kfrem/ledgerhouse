"""Build the LedgerHouse user, tester, maintenance and developer manual."""
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "manuals" / "LedgerHouse_User_Test_Maintenance_Developer_Manual_2026_08_03.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(88, 104, 122)
PALE_BLUE = "E8EEF5"
PALE_YELLOW = "FFF2CC"
WHITE = RGBColor(255, 255, 255)


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for index, width in enumerate(widths):
            row.cells[index].width = width
            for cell in (row.cells[index],):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(cell)


def mark_header_row(row):
    """Mark a Word table row as a repeating semantic header."""
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def configure_styles(document):
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in (
        ("Title", 24, RGBColor(0, 0, 0), 0, 6),
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = style_name != "Normal"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    if "Manual Note" not in styles:
        note = styles.add_style("Manual Note", WD_STYLE_TYPE.PARAGRAPH)
        note.font.name = "Calibri"
        note.font.size = Pt(10)
        note.font.color.rgb = DARK_BLUE
        note.paragraph_format.space_before = Pt(4)
        note.paragraph_format.space_after = Pt(6)
        note.paragraph_format.left_indent = Inches(0.15)


def add_header_footer(document):
    section = document.sections[0]
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("LedgerHouse | Test and Operations Manual")
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("LedgerHouse confidential test manual | ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def add_paragraph(document, text, bold_prefix=None, style=None):
    paragraph = document.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        run = paragraph.add_run(bold_prefix)
        run.bold = True
        paragraph.add_run(text[len(bold_prefix):])
    else:
        paragraph.add_run(text)
    return paragraph


def add_bullets(document, items):
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def add_steps(document, items):
    for item in items:
        document.add_paragraph(item, style="List Number")


def add_callout(document, title, text, warning=False):
    table = document.add_table(rows=1, cols=1)
    mark_header_row(table.rows[0])
    table.autofit = False
    table.cell(0, 0).width = Inches(6.5)
    cell = table.cell(0, 0)
    set_cell_shading(cell, PALE_YELLOW if warning else PALE_BLUE)
    set_cell_margins(cell, top=110, bottom=110, start=150, end=150)
    paragraph = cell.paragraphs[0]
    label = paragraph.add_run(title + "\n")
    label.bold = True
    label.font.color.rgb = DARK_BLUE
    paragraph.add_run(text)
    document.add_paragraph()


def add_table(document, headings, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headings))
    table.style = "Table Grid"
    mark_header_row(table.rows[0])
    header = table.rows[0].cells
    for index, heading in enumerate(headings):
        set_cell_shading(header[index], PALE_BLUE)
        run = header[index].paragraphs[0].add_run(heading)
        run.bold = True
        run.font.color.rgb = DARK_BLUE
    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            cells[index].text = value
    if widths:
        set_table_widths(table, widths)
    document.add_paragraph()
    return table


def add_section_heading(document, title, subtitle=None):
    document.add_heading(title, level=1)
    if subtitle:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(subtitle)
        run.italic = True
        run.font.color.rgb = MUTED


def build_document():
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    configure_styles(document)
    add_header_footer(document)

    # Cover
    document.add_paragraph()
    document.add_paragraph()
    brand = document.add_paragraph()
    brand.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = brand.add_run("LEDGERHOUSE")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(15)
    run.font.color.rgb = BLUE
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("User, Test, Maintenance and Developer Manual")
    run.bold = True
    run.font.size = Pt(24)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Controlled pre-market testing edition | 3 August 2026")
    run.font.size = Pt(12)
    run.font.color.rgb = MUTED
    document.add_paragraph()
    add_callout(
        document,
        "Purpose",
        "This handbook explains what LedgerHouse is for, how each role uses it, how to test it safely, how to maintain it and how future developers should extend it. It is written for people with no accounting-software experience as well as technical staff.",
    )
    add_table(document, ["Document owner", "Use"], [
        ("Product owner / test coordinator", "Issue to everyone involved in controlled testing."),
        ("Practice / accountant", "Use the practice workflows and sign-off checks."),
        ("Client tester", "Use the client portal workflow and feedback questions."),
        ("Maintenance staff", "Use the service, backup, incident and release sections."),
        ("Future developer", "Use the architecture, coding and verification sections."),
    ], [Inches(1.85), Inches(4.65)])
    document.add_page_break()

    add_section_heading(document, "1. What LedgerHouse Is", "A simple explanation for everyone")
    add_paragraph(document, "LedgerHouse is an accountant-supervised accounts-department platform for small and medium businesses. It is designed to collect business records, organise the work needed to turn those records into reliable books, help a practice or internal finance team review the work, and give the client clear reports and approval requests.")
    add_paragraph(document, "It is not a replacement for professional judgement. The client supplies records and decisions; the accountant or authorised finance user checks the evidence, controls the bookkeeping process and explains the outcome.")
    add_table(document, ["LedgerHouse does", "LedgerHouse does not do by itself"], [
        ("Stores and organises uploaded evidence, bank activity, journals, reports, tasks, client questions and approvals.", "Decide whether an expense is allowable, give tax advice or replace an accountant's review."),
        ("Produces management reports, CSV/PDF packs, controlled accounts-production records and local iXBRL preparation.", "Prove that a government filing has been accepted merely because a document was generated."),
        ("Keeps tenant data separated and records audit events for controlled work.", "Use a shared HMRC or Companies House application belonging to another product."),
    ], [Inches(3.25), Inches(3.25)])
    add_callout(document, "Testing rule", "Use synthetic or expressly authorised test data during this test phase. Do not upload client passwords, government access codes, payment-card details or any other secret into a test record.", True)

    # Keep this full role map together so printed copies retain the column headings.
    document.add_page_break()
    add_section_heading(document, "2. Roles and Access")
    add_table(document, ["Role", "Main purpose", "Main areas"], [
        ("Client user", "Provide records and understand the business position.", "Client portal, uploads, questions, approvals, reports."),
        ("Client approver", "Make formal client decisions.", "Approval packs, VAT/accounting decisions, client requests."),
        ("Practice bookkeeper", "Review incoming evidence and organise routine bookkeeping.", "Evidence review, banking review, ledger review, work queues."),
        ("Practice accountant / manager", "Approve work, supervise exceptions and prepare reporting / statutory work.", "Practice dashboard, accounts department, accounts production, reports."),
        ("System administrator", "Manage access, environment and safe releases.", "Django admin, deployment, backups, security settings."),
        ("Test coordinator", "Run the test plan and log evidence of outcomes.", "All areas, test checklist, defect log."),
        ("Future developer", "Improve the product without breaking accounting controls.", "Source code, migrations, tests, local Docker stack."),
    ], [Inches(1.35), Inches(2.15), Inches(3.0)])
    add_paragraph(document, "Practice pages are restricted to staff users. Client users should only see their own organisation's information. A client tester must report immediately if another client's name, document, bank line or report is visible.")

    add_section_heading(document, "3. Before Anyone Starts Testing")
    add_steps(document, [
        "The test coordinator creates or selects a test client company. Use a realistic but synthetic company name and records.",
        "Give each tester a separate account where possible. Do not share the production administrator account.",
        "Agree the test period, test transactions and expected outcome before uploading records.",
        "Record the tester, date, route used, action taken, expected result and actual result in the defect or test log.",
        "Classify any issue as blocked, defect, usability issue, data issue or training question. Do not silently work around it.",
    ])
    add_callout(document, "Local trial account", "The development environment may include a simple local trial account. It exists only for local testing and must never be used in a public deployment. The test coordinator should issue test access separately.", True)

    add_section_heading(document, "4. Client User Manual")
    document.add_heading("4.1 The client portal", level=2)
    add_paragraph(document, "The client portal is the simple side of LedgerHouse. It is for people who want to know what their finance team needs, what is being processed, what needs approval and what the latest figures mean.")
    add_bullets(document, [
        "Start with the overview: it should show the selected business, outstanding requests, key account-health items and available reports.",
        "Use the navigation to move to reports, approvals or questions rather than trying to edit technical accounting records.",
        "Read warnings as requests for attention, not as evidence that the books are wrong. Ask the practice team when the reason is unclear.",
    ])
    document.add_heading("4.2 Uploading documents", level=2)
    add_steps(document, [
        "Choose the correct client company before uploading anything.",
        "Upload one clear document at a time: invoice, receipt, bank file, expense image, spreadsheet or supporting document.",
        "Add a short description that tells the reviewer what the document is for, for example 'June office rent invoice'.",
        "Check the upload confirmation and return to the evidence list if you need to confirm it is present.",
        "Do not alter or delete a document after it has been used in accounting work. Ask the practice team to correct it through the controlled review process.",
    ])
    add_paragraph(document, "Expected outcome: the evidence is visible to the authorised practice team and is marked for review. An upload does not automatically mean it has been posted to the books.")
    document.add_heading("4.3 Questions and approvals", level=2)
    add_bullets(document, [
        "Use a client request when you need to explain a transaction, provide a missing document or ask a finance question.",
        "Read the summary before approving. Approval means you are confirming the stated item, not approving every accounting or tax consequence outside that summary.",
        "Add a note when approving or rejecting if the decision needs context.",
        "Once a decision is recorded, contact the practice team rather than attempting to reverse it yourself.",
    ])
    document.add_heading("4.4 Reports and downloads", level=2)
    add_paragraph(document, "Reports are designed to explain the current figures in plain English. Use the management report for explanation, CSV where analysis is needed in a spreadsheet, PDF for a fixed record, and the management pack for a fuller review. Compare the period and company name before relying on any result.")

    add_section_heading(document, "5. Practice Bookkeeper Manual")
    document.add_heading("5.1 Daily start", level=2)
    add_steps(document, [
        "Open Practice and select the client company you are working on.",
        "Review the work queues: evidence requiring review, unmatched bank lines, journals requiring review, client requests and approvals.",
        "Work the highest-risk or oldest item first. Do not post a journal just to clear a queue.",
        "Open the source document before changing its review status.",
    ])
    document.add_heading("5.2 Evidence review", level=2)
    add_bullets(document, [
        "Check the file is readable, complete and belongs to the selected client.",
        "Check supplier/customer, date, gross amount, VAT amount and suggested account code against the document.",
        "Use a review note to explain uncertainty, missing information or the reason for a coding decision.",
        "Use 'Needs information' when evidence is insufficient. Do not guess a material amount or tax treatment.",
        "Move an item to ready for posting only after its evidence is adequate and the intended treatment is clear.",
    ])
    document.add_heading("5.3 Bank and ledger review", level=2)
    add_bullets(document, [
        "Review unmatched bank lines against invoices, bills and supporting evidence.",
        "Record the reason for a match or leave it for further review. Do not force a match to make a bank balance look tidy.",
        "Post journals only when the journal is balanced and supported by evidence or an approved schedule.",
        "Treat a closed period or VAT-locked period as a control. Escalate required corrections to an authorised reviewer.",
    ])
    add_callout(document, "Stop and escalate", "Stop work and tell the practice manager if the document belongs to a different client, a bank movement cannot be explained, a material figure is missing, a duplicate is suspected, an approval is unclear or a locked period needs changing.", True)

    add_section_heading(document, "6. Practice Accountant and Manager Manual")
    document.add_heading("6.1 Supervision and exception register", level=2)
    add_paragraph(document, "The practice dashboard and exception register are the control centre. Use them to see what is waiting, who owns it and what prevents month end, reporting or client release.")
    add_steps(document, [
        "Review the exception register at the start of the day and before month end.",
        "Assign an owner, status and due date where the task is actionable.",
        "Resolve or document each blocker: missing document, unmatched bank line, unreviewed evidence, approval pending or month-close issue.",
        "Review audit events before overriding a workflow decision or releasing a significant report.",
        "Use client requests to obtain missing information; do not replace an evidence gap with an unsupported journal.",
    ])
    document.add_heading("6.2 Month end and recurring accounting", level=2)
    add_bullets(document, [
        "Check bank reconciliation, evidence review and outstanding exceptions before producing a management pack.",
        "Review recurring schedules such as accruals, prepayments and fixed-asset depreciation before postings are approved.",
        "Check the trial balance and report action points. A balanced trial balance does not by itself prove the underlying treatment is correct.",
        "Use the month-end close process only after the open exceptions and approval requirements have been dealt with or expressly carried forward.",
    ])
    # Avoid leaving this operationally important subsection heading alone at a page foot.
    document.add_page_break()
    document.add_heading("6.3 Accounts production and statutory readiness", level=2)
    add_paragraph(document, "Accounts Production is a controlled register of drafts, approvals and release history. It supports Micro-entity, Small unaudited full, Small unaudited filleted and Dormant accounts in the current declared scope. Amended accounts are variants of those types.")
    add_bullets(document, [
        "Confirm the correct accounts type, company number, period and client approval before preparing a filing package.",
        "Use the iXBRL preview and readiness messages as preflight checks. They are not provider acceptance.",
        "A gateway acknowledgement means the package was received for processing; it is not an accepted filing. Wait for a confirmed provider status.",
        "Do not enable test or live sending from a routine user interface. This is a controlled release-manager action.",
    ])

    add_section_heading(document, "7. Test Coordinator Manual")
    add_paragraph(document, "The test coordinator is responsible for proving that the workflow is understandable and safe for real people. A test is complete only when the expected result, actual result and evidence are recorded.")
    add_table(document, ["Test journey", "Minimum steps", "Pass condition"], [
        ("Client intake", "Log in, select test company, upload a clear invoice, add description.", "Evidence appears for the correct client and is awaiting review."),
        ("Evidence review", "Practice user opens the file, checks values, adds note and changes review status.", "Review state and audit detail are visible; no other client data appears."),
        ("Bank review", "Review an unmatched test bank line and either match it or mark it as needing information.", "No forced match; status and reason are retained."),
        ("Report release", "Open report, CSV, PDF and management pack for the selected client.", "Correct company/period, readable output, no cross-client data."),
        ("Approval", "Create or open an approval, approve it with a note, then re-open it.", "Decision, user and time are retained; repeat decision is controlled."),
    ], [Inches(1.45), Inches(3.1), Inches(1.95)])
    document.add_heading("7.0.1 Security and provider checks", level=2)
    add_table(document, ["Test journey", "Minimum steps", "Pass condition"], [
        ("Access control", "Use a client account to attempt a practice route.", "Access is denied or redirected; client cannot open practice work."),
        ("Government safety", "Open HMRC and Companies House areas without configured live credentials.", "Screen clearly shows configuration/blocker state; no external filing is claimed."),
    ], [Inches(1.45), Inches(3.1), Inches(1.95)])
    # The defect template must begin with its own heading, not as a continuation
    # beneath the larger journey matrix.
    document.add_page_break()
    document.add_heading("7.1 Defect report template", level=2)
    add_table(document, ["Field", "What to record"], [
        ("Reference", "A unique identifier, for example LH-TEST-001."),
        ("Role and route", "Who was testing and the exact page or menu route."),
        ("Steps", "Every click/type/upload action needed to reproduce the issue."),
        ("Expected result", "What the manual says should happen."),
        ("Actual result", "What happened, including visible error text."),
        ("Evidence", "Screenshot without secrets, test document name, time and browser/device."),
        ("Severity", "Blocker, high, medium, low or usability suggestion."),
    ], [Inches(1.7), Inches(4.8)])

    add_section_heading(document, "8. System Administrator and Maintenance Manual")
    document.add_heading("8.1 Day-to-day service checks", level=2)
    add_steps(document, [
        "Confirm the web and database containers are running: `docker compose ps`.",
        "Check application health: `curl http://127.0.0.1:8000/healthz/` locally, or the public HTTPS health URL after deployment.",
        "Review container logs for errors: `docker compose logs --tail=100 web`.",
        "Run the Django check after configuration changes: `docker compose exec -T web python manage.py check`.",
        "Investigate unavailable health responses before allowing users to continue accounting work.",
    ])
    document.add_heading("8.2 Backups and restore rehearsal", level=2)
    add_paragraph(document, "The database contains the ledger and evidence metadata. Backups must therefore be monitored as an operational responsibility, not treated as a one-time setup task.")
    add_bullets(document, [
        "Take scheduled provider/database backups in production and protect them separately from the application server.",
        "Run a local checksummed backup and restore rehearsal after material changes: `docker compose exec -T web python manage.py backup_local_database --restore-rehearsal`.",
        "Record the date, operator, backup location, result and any recovery-time observation.",
        "Never test a restore by overwriting a live database. Use a controlled copy or rehearsal path.",
    ])
    document.add_heading("8.3 Incidents", level=2)
    add_table(document, ["Situation", "Immediate action", "Do not do"], [
        ("Health endpoint is unavailable", "Stop deployment activity, inspect web/database logs and preserve evidence.", "Do not restart repeatedly without recording the failure."),
        ("Suspected cross-client access", "Remove access, preserve audit evidence and escalate to the product owner immediately.", "Do not alter or delete audit records."),
        ("Provider rejection", "Keep the response, review the correction action and send a new controlled submission only when authorised.", "Do not mark a filing accepted or reuse a rejected submission number."),
        ("Credential exposure", "Revoke or rotate through the provider, remove the exposure and document the incident.", "Do not paste the secret into tickets, chat, code or reports."),
    ], [Inches(1.55), Inches(2.75), Inches(2.2)])

    add_section_heading(document, "9. Release Manager Manual")
    add_paragraph(document, "A public deployment is a controlled change. The release manager owns the go/no-go decision and must distinguish local test evidence from proof that the public server and providers are working.")
    document.add_heading("9.1 Pre-release checklist", level=2)
    add_bullets(document, [
        "All automated tests pass from the current checkout.",
        "No migration drift: `python manage.py makemigrations --check --dry-run`.",
        "Production Compose resolves with production-only environment values.",
        "A real secret key, Fernet field-encryption key and database passwords exist outside Git.",
        "The restricted database role is in use for the web service, not a superuser or BYPASSRLS role.",
        "DNS points to the server and Caddy has obtained a valid certificate for the public domain.",
        "The public `/healthz/` endpoint and an authenticated client/practice smoke test pass.",
        "Backup, restore, incident contact and support route are agreed.",
    ])
    document.add_heading("9.2 Deployment sequence", level=2)
    add_steps(document, [
        "Place the production-only `.env` file on the server. Do not copy it into source control.",
        "Install the LedgerHouse Caddy site configuration and validate it before reloading Caddy.",
        "Run `deploy/hetzner/deploy_ledgerhouse.sh` from the server checkout.",
        "Confirm HTTPS, `/healthz/`, login, a client route, a practice route and a report download on the public domain.",
        "Record the deployed version, time, operator, health result and rollback point.",
    ])
    add_callout(document, "Provider release controls", "HMRC and Companies House sending must remain disabled until the separately-approved LedgerHouse credentials, scope, user authority and controlled provider test are confirmed. A generated file, an OAuth screen or a gateway acknowledgement is not final acceptance.", True)

    # The architecture map is a reference table; start it on a clean page.
    document.add_page_break()
    add_section_heading(document, "10. Future Developer Manual")
    document.add_heading("10.1 Architecture map", level=2)
    add_table(document, ["Area", "Location", "Developer rule"], [
        ("Web routes", "`ledgerhouse/urls.py` and `accounting/views.py`", "Protect practice routes and preserve tenant selection/access checks."),
        ("Accounting domain", "`accounting/models.py`, operations, reports and workflow modules", "Keep accounting entries balanced, auditable and tenant-scoped."),
        ("Government adapters", "`accounting/mtd.py` and `accounting/companies_house.py`", "Fail closed; never fabricate provider IDs, acceptance or a filing receipt."),
        ("Database integrity", "PostgreSQL migrations and database checks", "Do not replace PostgreSQL controls with SQLite-only or application-only logic."),
        ("User interface", "`templates/accounting/` and `static/accounting/`", "Keep client wording plain and practice workflows dense but understandable."),
        ("Deployment", "Docker, `deploy/`, `DEPLOYMENT.md`", "Keep secrets environment-only and test production configuration before release."),
    ], [Inches(1.25), Inches(2.35), Inches(2.9)])
    document.add_heading("10.2 Change process", level=2)
    add_steps(document, [
        "Read the affected route, model, migrations and existing tests before editing.",
        "Check `git status` and preserve work already in progress.",
        "Add focused tests for a new rule, including failure and permission cases.",
        "Run focused tests, fix failures and then run the full Django test suite.",
        "Run `makemigrations --check --dry-run`, the local acceptance suite and the relevant rehearsal after meaningful workflow changes.",
        "Browser-check any changed user-facing route at desktop and narrow mobile width.",
        "Update this manual, the release report and deployment documentation when behaviour or operations change.",
    ])
    document.add_heading("10.3 Non-negotiable technical rules", level=2)
    add_bullets(document, [
        "Never store secrets, access codes, tokens or raw provider responses in Git, screenshots or test reports.",
        "Never create a provider success, receipt, correlation ID or accepted filing from a mock response.",
        "Never reuse another product's HMRC client ID, secret, redirect URI, subscription, presenter credentials or package reference.",
        "Do not bypass tenant isolation, audit records, period locks or VAT locks to make a test easier.",
        "Do not change a database record directly in production when a controlled workflow or audited correction is available.",
    ])

    add_section_heading(document, "11. HMRC and Companies House Guidance")
    document.add_heading("11.1 HMRC", level=2)
    add_paragraph(document, "The HMRC area is a sandbox and integration status workspace. It can show whether LedgerHouse's own configuration is complete and provide the controlled VAT workflow. It must not use credentials or subscriptions from another product.")
    add_bullets(document, [
        "Use a separate LedgerHouse application, client secret, redirect URI and API subscription.",
        "Test harmless sandbox endpoints first, then controlled OAuth and VAT journeys only when proper user authority and fraud-prevention data are available.",
        "A successful login or OAuth page does not prove that a filing succeeded.",
        "Production calls require production approval, production credentials, real vendor information and an authorised user journey.",
    ])
    document.add_heading("11.2 Companies House", level=2)
    add_paragraph(document, "The Companies House workspace prepares and controls accounts-filing packages in the declared accounts scope. It records submission sequencing, sanitised evidence, validation messages and status polling. It does not declare a filing accepted until the gateway provides a recognised accepted status.")
    add_bullets(document, [
        "Test presenter credentials and test package references are for the LedgerHouse product only.",
        "Submission numbers must be unique and incremental for that product. Do not reuse a number after a rejection or failed build.",
        "Gateway acknowledgement is pending, not acceptance. Poll and record the resulting status.",
        "Live filing stays locked until the provider has approved the relevant scope and live configuration has been explicitly enabled.",
    ])

    add_section_heading(document, "12. Troubleshooting and Glossary")
    document.add_heading("12.1 Common questions", level=2)
    add_table(document, ["Question", "Answer"], [
        ("Why is a document not in the reports yet?", "Uploading puts evidence into review. It may still need coding, bank matching, approval or posting."),
        ("Why can I not post into a period?", "A closed period or VAT lock is protecting the audit trail. Escalate the correction to an authorised reviewer."),
        ("Why is a government button disabled?", "LedgerHouse is deliberately blocking an incomplete or unapproved provider configuration. It is safer than sending a request with the wrong credentials."),
        ("Why does a filing show Pending?", "The provider has not returned a confirmed final outcome. Pending is not accepted."),
        ("What is an audit trail?", "A dated record of an important action, who performed it and the controlled workflow it affected."),
        ("What is a tenant?", "One client company and its separate LedgerHouse data boundary."),
    ], [Inches(2.1), Inches(4.4)])
    # Keep the glossary as a single searchable, printable reference table.
    document.add_page_break()
    document.add_heading("12.2 Glossary", level=2)
    add_table(document, ["Term", "Plain-English meaning"], [
        ("Evidence", "The supporting document or file behind an accounting entry."),
        ("Journal", "A balanced accounting record that moves values between accounts."),
        ("Reconciliation", "Checking bank activity against the accounting records and supporting evidence."),
        ("Management pack", "A structured report for owners/managers explaining performance and actions."),
        ("iXBRL", "A web-readable accounts document format used for some statutory filing processes."),
        ("Sandbox", "A provider test environment. It is not the live service."),
        ("Fail closed", "Stop a risky action when evidence or configuration is incomplete, instead of guessing or pretending it worked."),
    ], [Inches(1.55), Inches(4.95)])

    # Preserve the blank sign-off rows as one printable sheet.
    document.add_page_break()
    add_section_heading(document, "13. Final Test Sign-off")
    add_paragraph(document, "The test coordinator should complete this page only after the required journeys have been run with the assigned users and issues have been logged or resolved.")
    add_table(document, ["Area", "Tester", "Date", "Result / reference"], [
        ("Client portal and document upload", "", "", ""),
        ("Practice evidence and banking review", "", "", ""),
        ("Ledger review, period controls and reports", "", "", ""),
        ("Client requests and approvals", "", "", ""),
        ("Accounts Production and iXBRL preview", "", "", ""),
        ("Access control / tenant separation", "", "", ""),
        ("Maintenance health and backup rehearsal", "", "", ""),
        ("Public deployment smoke test", "", "", ""),
    ], [Inches(2.2), Inches(1.3), Inches(1.1), Inches(2.0)])
    add_callout(document, "Sign-off statement", "Signing this test record means the listed journeys were executed and the results are recorded. It does not override external provider approval, legal review or a production release decision.", True)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.core_properties.title = "LedgerHouse User, Test, Maintenance and Developer Manual"
    document.core_properties.subject = "Controlled pre-market testing edition"
    document.core_properties.author = "LedgerHouse"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
