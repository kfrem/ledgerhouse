from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "deployment"
OUT_FILE = OUT_DIR / "ledgerhouse_server_provider_brief_2026_08_05.docx"

NAVY = "17324D"
TEAL = "147D8C"
LIGHT = "EAF3F5"
GREY = "F2F4F5"
DARK = "1E2933"


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    set_table_widths(table, widths)
    header_cells = table.rows[0].cells
    set_repeat_table_header(table.rows[0])
    for cell, header in zip(header_cells, headers):
        set_cell_fill(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        run = paragraph.add_run(header)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            if row_index % 2 == 1:
                set_cell_fill(cell, GREY)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            run.font.name = "Arial"
            run.font.size = Pt(8.7)
    doc.add_paragraph()
    return table


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph()
    paragraph.style = f"Heading {level}"
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.color.rgb = RGBColor.from_string(NAVY if level == 1 else TEAL)
    return paragraph


def add_body(doc, text, bold_lead=None):
    paragraph = doc.add_paragraph()
    paragraph.style = "Normal"
    paragraph.paragraph_format.space_after = Pt(5)
    if bold_lead:
        lead = paragraph.add_run(bold_lead)
        lead.bold = True
    paragraph.add_run(text)
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(item)
        run.font.name = "Arial"
        run.font.size = Pt(9.4)


def add_numbered(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(item)
        run.font.name = "Arial"
        run.font.size = Pt(9.4)


def build_document():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    for name, size in (("Heading 1", 14), ("Heading 2", 11), ("Heading 3", 10)):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(NAVY if name == "Heading 1" else TEAL)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("LEDGERHOUSE | SERVER PROVIDER BRIEF")
    run.font.name = "Arial"
    run.font.bold = True
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(TEAL)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Confidential infrastructure request | 5 August 2026")
    footer_run.font.name = "Arial"
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor.from_string("607080")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run("LedgerHouse")
    run.font.name = "Arial"
    run.font.bold = True
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor.from_string(NAVY)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run("UK production server and go-live requirements")
    run.font.name = "Arial"
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor.from_string(TEAL)

    callout = doc.add_table(rows=1, cols=1)
    callout.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = callout.cell(0, 0)
    set_cell_fill(cell, LIGHT)
    set_cell_margins(cell, top=130, start=150, bottom=130, end=150)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    lead = paragraph.add_run("Request to provider: ")
    lead.bold = True
    lead.font.name = "Arial"
    lead.font.color.rgb = RGBColor.from_string(NAVY)
    run = paragraph.add_run(
        "Provide a secure, managed or self-managed Linux VPS suitable for a Docker-based, multi-tenant UK accounting platform. "
        "This is not a shared-hosting request. The platform requires a persistent PostgreSQL database, HTTPS, backups, firewall control, SSH-key access and a fixed public IPv4 address."
    )
    run.font.name = "Arial"
    doc.add_paragraph()

    add_heading(doc, "1. Product Summary")
    add_body(
        doc,
        "LedgerHouse is a SaaS-style full accounts department platform for UK small and medium businesses and accounting practices. Clients submit invoices, receipts, bank files, photos and spreadsheets. Practice staff review the evidence, complete bookkeeping and controls, produce management information, seek client approval and prepare statutory-accounting and tax workflows.",
    )
    add_body(
        doc,
        "It is designed as a multi-tenant web application: each client organisation's data is logically segregated, with role-based client and practice workflows, immutable audit events and PostgreSQL-backed accounting controls.",
    )

    add_heading(doc, "2. Current Product Functions")
    add_table(doc, ["Area", "LedgerHouse capability"], [
        ("Client portal", "Secure sign-in, document/evidence upload, visible practice progress, questions, approvals and report-pack access."),
        ("Evidence intake", "Invoice, receipt, statement and spreadsheet intake; document categorisation, validation, review queues and clear exception outcomes."),
        ("Bookkeeping", "Sales and purchase workflow, double-entry ledger support, bank transaction review, matching and reconciliation controls."),
        ("Accounting operations", "Recurring accruals, prepayments, depreciation schedules, period-close safeguards, payroll summaries and corporation-tax estimates."),
        ("Management information", "Management reports, financial-position reporting, client approval packs, release history and downloadable reporting outputs."),
        ("Practice workspace", "Client work queues, assignments, controlled status changes, evidence review, exception register and auditable operational timeline."),
        ("Statutory accounts", "Local generation and validation of report-style XHTML/iXBRL accounts for the supported Companies House accounts filing types."),
        ("Government integration", "Companies House public-data integration and XML Gateway filing preparation. HMRC integration configuration is isolated to LedgerHouse and remains provider-authorisation dependent."),
        ("Security and controls", "Tenant controls, PostgreSQL row-level-security design, audit records, encryption-key configuration, client approval gates and fail-closed filing controls."),
    ], [1.6, 5.55])

    add_heading(doc, "3. Technology Architecture")
    add_table(doc, ["Component", "Required technology"], [
        ("Application", "Python Django application served by Gunicorn in a Docker container."),
        ("Database", "PostgreSQL 16 in a separate Docker service with persistent storage. PostgreSQL is mandatory for accounting constraints, closed-period controls, audit protections and tenant isolation."),
        ("Reverse proxy", "Caddy on the host, terminating TLS and forwarding HTTPS traffic to the internal application container."),
        ("Deployment", "Docker Engine and Docker Compose v2; Git-based application deployment and environment-only configuration."),
        ("Security", "HTTPS, secure cookies, HSTS, SSH keys only, firewall, fail2ban, unattended security updates and a non-superuser application database role after initial migration."),
        ("Health monitoring", "Application health endpoint at `/healthz/`; Docker health checks for the database and web service."),
    ], [1.6, 5.55])

    add_heading(doc, "4. Server Specification Requested")
    add_body(doc, "The provider should quote the recommended specification. The minimum specification is acceptable only for a tightly controlled early pilot with low concurrent use.")
    add_table(doc, ["Requirement", "Early pilot minimum", "Recommended for UK client launch"], [
        ("Virtual server", "2 vCPU", "4 vCPU"),
        ("Memory", "4 GB RAM", "8 GB RAM"),
        ("Storage", "40 GB NVMe SSD", "80 GB NVMe SSD, expandable"),
        ("Operating system", "Ubuntu Server 24.04 LTS", "Ubuntu Server 24.04 LTS"),
        ("Public networking", "1 static public IPv4; IPv6 desirable", "1 static public IPv4; IPv6 desirable"),
        ("Network transfer", "At least 1 TB/month or unmetered", "At least 2 TB/month or unmetered"),
        ("Location", "UK or EEA data centre", "UK data centre preferred; EEA acceptable with suitable UK GDPR terms"),
        ("Backups", "Daily backup and restore capability", "Daily encrypted off-host backup plus retained snapshots and documented restore procedure"),
        ("Availability", "Provider monitoring and support", "24/7 host monitoring, incident contact route and a published uptime commitment"),
    ], [1.45, 2.75, 2.95])

    add_heading(doc, "5. Non-Negotiable Provider Capabilities")
    add_bullets(doc, [
        "Full administrative control of the VPS, including SSH key access and permission to run Docker Engine and Docker Compose v2.",
        "A persistent disk/volume for PostgreSQL data. The database must not be ephemeral and must not be exposed directly to the public internet.",
        "Control of firewall rules or an equivalent provider firewall. Public inbound traffic is limited to TCP 80 and 443; SSH TCP 22 is restricted to the administrator's approved IP address where feasible.",
        "A fixed public IPv4 address for DNS, TLS certificate issuance and future government/provider allow-listing where required.",
        "Ability to create DNS A/AAAA records or clear delegation instructions for `ledgerhouse.finaccord.pro`.",
        "No forced control panel, application proxy or managed database layer that prevents Caddy, Docker, PostgreSQL volumes or standard Linux administration.",
        "A current data-processing agreement and clear statement of data-centre location, sub-processors, backup location and incident-notification process suitable for UK business data.",
    ])

    add_heading(doc, "6. Required Production Topology")
    add_body(doc, "The provider need only supply the server and network capabilities. LedgerHouse deploys its own application stack as follows:")
    add_numbered(doc, [
        "DNS directs `ledgerhouse.finaccord.pro` to the server's static IP address.",
        "Caddy listens publicly on HTTPS and obtains/renews TLS certificates automatically after DNS is valid.",
        "Caddy reverse-proxies requests only to the LedgerHouse Docker web container on the private host interface.",
        "The Django/Gunicorn web container connects to the PostgreSQL 16 container over the Docker network.",
        "PostgreSQL uses an encrypted-provider disk or persistent volume; its port is not published publicly.",
        "A separate private server directory holds environment configuration and the LedgerHouse-only Companies House credential file. It is mounted read-only into the application container and is never committed to source control.",
        "Backups run from the database volume to a separate encrypted, off-host location. Restore tests are performed before commercial onboarding.",
    ])

    add_heading(doc, "7. Security, Privacy and Operational Requirements")
    add_table(doc, ["Control", "Requirement"], [
        ("TLS", "TLS 1.2+ over HTTPS, automatic certificate renewal, HTTP-to-HTTPS redirect, secure cookies and HSTS enabled after certificate verification."),
        ("Access", "SSH key authentication only; no password login. Named administrative access and a documented offboarding procedure."),
        ("Firewall", "Inbound: 80/443 public; 22 restricted where possible. PostgreSQL and Docker service ports must not be publicly exposed."),
        ("Secrets", "Production secrets only in an ignored server-side environment file/private secret directory. Never in Git, tickets, chat, screenshots or container images."),
        ("Database", "Use the restricted LedgerHouse runtime database role after migration, not a PostgreSQL superuser. Back up daily and test restoration."),
        ("Monitoring", "Monitor host resource use, Docker container health, HTTPS certificate status, disk space, database backup success and `/healthz/`. Alert the operations contact on failure."),
        ("Logging", "Retain application and access logs for investigation without logging passwords, credentials, tokens or client financial documents."),
        ("Patch management", "Apply unattended security updates and scheduled maintenance/reboots. Notify the owner of material host incidents or planned disruptive maintenance."),
    ], [1.55, 5.6])

    add_heading(doc, "8. Go-Live Work Plan")
    add_numbered(doc, [
        "Provision the VPS, provide its static IPv4, SSH-key access and data-centre/backup details.",
        "Create the DNS A record for `ledgerhouse.finaccord.pro` and allow it to propagate.",
        "Install or confirm Ubuntu 24.04 LTS, Docker Engine, Docker Compose v2, Caddy, UFW, fail2ban and unattended security updates.",
        "Deploy the LedgerHouse code to the private application directory and create the production environment file and private secret directory with owner-only permissions.",
        "Start PostgreSQL and LedgerHouse with the production Docker Compose file; apply migrations using the schema-owner connection; then configure the restricted runtime database role.",
        "Configure Caddy for `ledgerhouse.finaccord.pro`, validate the configuration and verify an issued HTTPS certificate.",
        "Set up daily off-host database backup, retain a defined restore window and complete one documented restoration test before onboarding live clients.",
        "Run the production smoke checks: database health, `/healthz/`, sign-in, client portal, practice workflow, uploads, reporting, audit logging and negative-access checks.",
        "Keep government filing transmission disabled until the deployment is verified and each client has supplied the necessary authority and approval. No testing is performed in the live Companies House service.",
    ])

    add_heading(doc, "9. Provider Acceptance Checklist")
    add_table(doc, ["Acceptance item", "Evidence requested"], [
        ("Server provisioned", "Hostname, static IP, operating-system version, CPU/RAM/storage and data-centre location."),
        ("Administrative access", "SSH key access confirmed; password login disabled or capable of being disabled."),
        ("Firewall", "Inbound-rule list showing only permitted services; no public PostgreSQL port."),
        ("Docker support", "Docker Engine and Compose v2 installed or permission confirmed to install and operate them."),
        ("Persistent storage", "Confirmation of durable storage and how disk expansion is requested."),
        ("Backup service", "Backup frequency, retention, encryption, storage location, restoration process, cost and any exclusions."),
        ("Monitoring/support", "Monitoring coverage, alert route, support hours, incident response targets and maintenance policy."),
        ("Privacy terms", "DPA, data-centre country, sub-processor list and breach/incident notification terms."),
    ], [2.05, 5.1])

    add_heading(doc, "10. Explicit Boundaries")
    add_body(doc, "A server provider supplies secure infrastructure. It does not itself grant authority to submit filings or access client bank/tax data. The following remain controlled application or client actions:")
    add_bullets(doc, [
        "HMRC production authorisation, user authority and API credentials.",
        "Companies House live filing only for supported document types, using the approved LedgerHouse package and an authorised client filing. Live service is not a test environment.",
        "Client Companies House authentication codes, client consent and final filing approval.",
        "Bank-feed provider consent, payment-provider account setup, email-provider setup and other third-party credentials.",
        "Legal, privacy and commercial terms review before broad public marketing or paid-client onboarding.",
    ])

    add_heading(doc, "11. Request for Quotation")
    add_body(doc, "Please respond with a quotation and confirmation against the following:")
    add_bullets(doc, [
        "Monthly price for the recommended 4 vCPU / 8 GB RAM / 80 GB NVMe specification, including VAT and all recurring charges.",
        "Separate cost for backups, snapshots, static IPv4, additional storage, bandwidth overages, managed support and server rebuilds.",
        "Data-centre location, backup location and data-processing terms.",
        "Whether root/administrator SSH access and Docker Compose are permitted.",
        "Provisioning lead time and the exact information needed to create the server today.",
        "The support route for an outage and the expected response times.",
    ])

    add_heading(doc, "12. Implementation Contact")
    add_body(doc, "Application: LedgerHouse, operated by KAFS Limited / KAFS Automation. Technical deployment will be performed by the LedgerHouse engineering team once the server, DNS and access details are available.")
    add_body(doc, "This brief deliberately excludes passwords, API tokens, presenter credentials, database passwords and client data. Providers must not request those values through ordinary email or support tickets.")

    doc.save(OUT_FILE)
    print(OUT_FILE)


if __name__ == "__main__":
    build_document()
